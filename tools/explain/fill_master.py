#!/usr/bin/env python3
"""Write the pipeline's results into the master spreadsheet (rows matched by the id in column P).

  tools/explain/.venv/bin/python tools/explain/fill_master.py "<Master.xlsx>"

  K  ready explanations (never overwrites a K the instructor already wrote)
  M  why a row needs the instructor, labelled: KEY WRONG? / NEW EXHIBIT / AUDIT / REVIEW /
     AMBIGUOUS / MISSING EXHIBIT; the parser's own note (community vote, low consensus) is kept
     after the label. A row whose blind answer matched the key with high confidence is ready,
     which also settles the parser's community-vote note, so M is cleared.
  N  proposed explanation, O proposed exhibit: copy into K / L to approve.

Close the workbook in Excel first; the caller backs it up.
"""
import os, re, sys
import openpyxl
from openpyxl.styles import Alignment, Font
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import check
from common import LETTERS, bank, letters, read

def nums(lets):
    return ','.join(str(LETTERS.index(x) + 1) for x in lets)

def audit_note(q, a):
    if not a:
        return ''
    probs = []
    if sorted(a['answer']) != letters(q['answer']):
        probs.append(f"blind audit answered {nums(a['answer'])}, key is {nums(letters(q['answer']))}")
    if not a['sufficient']:
        probs.append('exhibit lacks data needed for the answer')
    if not a['unique']:
        probs.append('more than one answer is defensible')
    if a['issues']:
        probs.append(a['issues'])
    if a['fix']:
        probs.append('Suggested fix: ' + a['fix'])
    return 'AUDIT (Fable, blind): ' + ' '.join(p.rstrip('.') + '.' for p in probs) if probs else ''

def review(q, r1, r2, ex, v):
    """-> (K, M, N, O) from the pipeline results for one question."""
    status, expl, exhibit, reasons = check.assess(q, r1, r2, ex, v)
    why = '; '.join(reasons)
    if status == 'ready':
        return expl, '', '', ''
    if status == 'NEW EXHIBIT':
        return '', ('NEW EXHIBIT: generated to replace one missing from the source; a blind check picked the '
                    'keyed answer. To publish, copy O into L and N into K.'), expl, exhibit
    if status == 'KEY WRONG?':
        better = r2['better_answer']
        m = (f"KEY WRONG? Suggested answer: {nums(better)} "
             f"({' + '.join(q['options'][LETTERS.index(x)] for x in better)}). {r2['note']}")
        n = ''
        if r1 and sorted(r1['answer']) == sorted(better):
            n = r1['explanation']
            m += ' Column N explains the suggested answer; if you change H, copy N into K.'
        return '', m, n, ''
    if status == 'EXHIBIT: drop':
        return '', f'MISSING EXHIBIT: could not be rebuilt ({why}). Delete this row or add an exhibit in L.', '', exhibit
    if status == 'EXHIBIT: pending':
        return '', f'MISSING EXHIBIT: not processed yet ({why}).', '', ''
    note = (r2 or {}).get('note') or (r1 or {}).get('note') or ''
    return '', f"{status.upper()}: {why}." + (f' {note}' if note else '') + (
        ' If N is right, copy it into K.' if expl else ''), expl, ''

def main():
    path = sys.argv[1]
    B = {q['id']: q for q in bank()}
    p1, p2, ex, ver, au = (read(n) for n in ('pass1.json', 'pass2.json', 'exhibits.json', 'verify.json', 'audit.json'))
    wb = openpyxl.load_workbook(path); ws = wb.active
    counts = {}
    for r in range(2, ws.max_row + 1):
        qid = ws.cell(r, 16).value
        q = B.get(qid)
        if not q or ws.cell(r, 11).value:                       # not in the bank, or K already written
            continue
        if qid not in p1 and qid not in ex:                      # not processed yet: leave the row alone
            continue
        k, m, n, o = review(q, p1.get(qid), p2.get(qid), ex.get(qid), ver.get(qid))
        a = audit_note(q, au.get(qid))
        if a:                                                    # an audit problem always goes to the instructor
            m = a + (' ' + m if m else '')
            n, k = n or k, ''
        parser_note = str(ws.cell(r, 13).value or '')
        if m and parser_note and not parser_note.startswith(m.split(':')[0]):
            m += ' | Parser: ' + parser_note
        label = (m.split(':')[0] if m else 'ready')
        counts[label] = counts.get(label, 0) + 1
        ws.cell(r, 11).value = k or None
        ws.cell(r, 13).value = m or None
        ws.cell(r, 14).value = n or None
        ws.cell(r, 15).value = o or None
        for c in (11, 13, 14, 15):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical='top')
        ws.cell(r, 15).font = Font(name='Consolas', size=10)
    wb.save(path)
    for label, c in sorted(counts.items(), key=lambda x: -x[1]):
        print(f'{c:>5}  {label}')

if __name__ == '__main__':
    main()
